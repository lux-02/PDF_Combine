import io

import pytest

from pdf_studio.conversion import DOCX_TEXT_SUPPORT, docx_to_pdf_simple, iter_docx_blocks


def make_minimal_docx_bytes() -> bytes:
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("Hello World")
    doc.add_paragraph("Second paragraph")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def make_docx_with_table() -> bytes:
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def make_docx_with_korean() -> bytes:
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("안녕하세요 한글 텍스트입니다")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


@pytest.mark.skipif(not DOCX_TEXT_SUPPORT, reason="python-docx or reportlab not available")
class TestDocxToPdfSimple:
    def test_returns_buffer_and_flags(self):
        docx_bytes = make_minimal_docx_bytes()
        buf, font_used, has_non_ascii = docx_to_pdf_simple(docx_bytes)
        assert hasattr(buf, "read")
        assert isinstance(font_used, bool)
        assert isinstance(has_non_ascii, bool)

    def test_output_is_valid_pdf(self):
        from pypdf import PdfReader
        docx_bytes = make_minimal_docx_bytes()
        buf, _, _ = docx_to_pdf_simple(docx_bytes)
        reader = PdfReader(buf)
        assert len(reader.pages) >= 1

    def test_ascii_only_has_no_non_ascii_flag(self):
        docx_bytes = make_minimal_docx_bytes()
        _, _, has_non_ascii = docx_to_pdf_simple(docx_bytes)
        assert has_non_ascii is False

    def test_korean_text_sets_non_ascii_flag(self):
        docx_bytes = make_docx_with_korean()
        _, _, has_non_ascii = docx_to_pdf_simple(docx_bytes)
        assert has_non_ascii is True

    def test_table_document_converts(self):
        from pypdf import PdfReader
        docx_bytes = make_docx_with_table()
        buf, _, _ = docx_to_pdf_simple(docx_bytes)
        reader = PdfReader(buf)
        assert len(reader.pages) >= 1

    def test_empty_document_returns_buffer(self):
        from docx import Document as DocxDocument
        doc = DocxDocument()
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        result_buf, font_used, has_non_ascii = docx_to_pdf_simple(buf.read())
        assert hasattr(result_buf, "read")
        assert isinstance(font_used, bool)
        assert has_non_ascii is False


@pytest.mark.skipif(not DOCX_TEXT_SUPPORT, reason="python-docx not available")
class TestIterDocxBlocks:
    def test_yields_paragraphs(self):
        from docx import Document as DocxDocument
        doc = DocxDocument()
        doc.add_paragraph("Para 1")
        doc.add_paragraph("Para 2")
        blocks = list(iter_docx_blocks(doc))
        texts = [b.text for b in blocks if hasattr(b, "text")]
        assert "Para 1" in texts
        assert "Para 2" in texts

    def test_yields_tables(self):
        from docx import Document as DocxDocument
        from docx.table import Table
        doc = DocxDocument()
        doc.add_table(rows=1, cols=2)
        blocks = list(iter_docx_blocks(doc))
        table_blocks = [b for b in blocks if isinstance(b, Table)]
        assert len(table_blocks) == 1

    def test_empty_document_yields_nothing(self):
        from docx import Document as DocxDocument
        doc = DocxDocument()
        for p in doc.paragraphs:
            p._element.getparent().remove(p._element)
        blocks = list(iter_docx_blocks(doc))
        assert blocks == []
